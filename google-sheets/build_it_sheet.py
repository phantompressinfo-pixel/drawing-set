"""The IT issue log, built as a real spreadsheet file.

Formatting cannot be applied through the Drive connector and the Apps Script
route was a dead end, so the styling is baked into an .xlsx instead: Google
keeps the fills, widths, frozen header and dropdowns when it converts.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NAVY, RULE, BAND = "FF022049", "FFD5DCE4", "FFF4F6F9"

STAFF = [("Date", 12), ("Your name", 16), ("Computer or device", 24),
         ("What is the problem?", 40),
         ("What were you doing when it happened?", 46),
         ("How urgent?", 20)]
IT    = [("Status", 16), ("IT notes", 40), ("Date fixed", 13)]

URGENCY  = '"Blocked - I cannot work,Slowing me down,Whenever you get a chance"'
STATUSES = '"New,In progress,Waiting on you,Done"'
LAST = 150

wb = Workbook(); ws = wb.active; ws.title = "Issues"

# Row 1 says who fills what, so the split is obvious without a legend.
ws["A1"] = ("REPORT AN IT ISSUE     Fill in columns A-F. "
            "Leave G-I to IT. One row per problem.")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws["A1"].font = Font(name="Montserrat", bold=True, size=11, color="FFFFFFFF")
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 34

cols = STAFF + IT
bottom = Border(bottom=Side(style="medium", color=NAVY))
for i, (title, width) in enumerate(cols, start=1):
    c = ws.cell(row=2, column=i, value=title)
    c.font = Font(name="Montserrat", bold=True, size=10,
                  color="FF6B7280" if i > len(STAFF) else NAVY)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = bottom
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[2].height = 30
ws.freeze_panes = "A3"

# The IT half gets a tint so nobody types in it by accident.
tint = PatternFill("solid", fgColor=BAND)
row_rule = Border(bottom=Side(style="thin", color=RULE))
for r in range(3, LAST + 1):
    for i in range(1, len(cols) + 1):
        c = ws.cell(row=r, column=i)
        c.font = Font(name="Montserrat", size=10)
        c.alignment = Alignment(vertical="center", wrap_text=i in (4, 5, 8))
        c.border = row_rule
        if i > len(STAFF):
            c.fill = tint

for rng, formula in ((f"F3:F{LAST}", URGENCY), (f"G3:G{LAST}", STATUSES)):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showDropDown = False          # openpyxl inverts this: False shows it
    ws.add_data_validation(dv); dv.add(rng)

wb.save("it-issue-log.xlsx")
print("written")
