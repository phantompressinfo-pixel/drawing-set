"""The Learning Sessions sign-up sheet, formatting baked into the file.

Fridays are pre-filled so nobody has to pick a date or add a row -- they
find an Open one, type their name in, and set Status to Claimed. Holiday
Fridays are left out rather than listed and crossed off.
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

NAVY, RULE = "FF022049", "FFD5DCE4"
OPEN, CLAIM, CONF, DONE = "FFF4F6F9", "FFFFF4D6", "FFE7F4EA", "FFEEEEEE"

COLS = [("Date", 16), ("Presenter", 20), ("Topic", 40), ("Format", 20),
        ("Notes for attendees", 38), ("Status", 14)]
FORMATS  = '"Demo,Walkthrough,Lessons learned,Project story,Guest speaker"'
STATUSES = '"Open,Claimed,Confirmed,Done"'

# Fridays the office is actually open. Thanksgiving Friday, Christmas and
# New Year's Day are skipped outright -- a listed date nobody can take is
# just a row people have to think about.
SKIP = {datetime.date(2026, 11, 27), datetime.date(2026, 12, 25),
        datetime.date(2027, 1, 1)}
d = datetime.date(2026, 9, 19)
d += datetime.timedelta(days=(4 - d.weekday()) % 7 or 7)
fridays = []
while len(fridays) < 20:
    if d not in SKIP:
        fridays.append(d)
    d += datetime.timedelta(days=7)

wb = Workbook(); ws = wb.active; ws.title = "Schedule"

ws["A1"] = ("LEARNING SESSIONS     Find a Friday marked Open, put your name "
            "and topic in, set Status to Claimed. That is the whole process.")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws["A1"].font = Font(name="Montserrat", bold=True, size=11, color="FFFFFFFF")
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 34

bottom = Border(bottom=Side(style="medium", color=NAVY))
for i, (title, width) in enumerate(COLS, start=1):
    c = ws.cell(row=2, column=i, value=title)
    c.font = Font(name="Montserrat", bold=True, size=10, color=NAVY)
    c.alignment = Alignment(vertical="center")
    c.border = bottom
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[2].height = 30
ws.freeze_panes = "A3"

row_rule = Border(bottom=Side(style="thin", color=RULE))
for n, day in enumerate(fridays):
    r = 3 + n
    ws.cell(row=r, column=1, value=day.strftime("%a %b %d, %Y"))
    ws.cell(row=r, column=6, value="Open")
    for i in range(1, 7):
        c = ws.cell(row=r, column=i)
        c.font = Font(name="Montserrat", size=10,
                      bold=(i == 1), color=NAVY if i == 1 else "FF000000")
        c.alignment = Alignment(vertical="center", wrap_text=(i == 5),
                                horizontal="center" if i in (4, 6) else "left")
        c.border = row_rule
    ws.row_dimensions[r].height = 26

last = 2 + len(fridays)
for rng, formula in ((f"D3:D{last}", FORMATS), (f"F3:F{last}", STATUSES)):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showDropDown = False          # openpyxl inverts this: False shows it
    ws.add_data_validation(dv); dv.add(rng)

# The row takes its colour from Status, so an open Friday is visible from
# across the room without anyone reading a legend.
band = f"A3:F{last}"
for value, fill in (("Open", OPEN), ("Claimed", CLAIM),
                    ("Confirmed", CONF), ("Done", DONE)):
    ws.conditional_formatting.add(band, FormulaRule(
        formula=[f'$F3="{value}"'],
        fill=PatternFill("solid", start_color=fill, end_color=fill)))

wb.save("learning-sessions-signup.xlsx")
print("rows:", len(fridays), "| first:", fridays[0], "| last:", fridays[-1])
